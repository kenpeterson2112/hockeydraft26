#!/usr/bin/env python3
"""import-projections.py — rebuild data/players.json from the projections workbook.

    python3 tools/import-projections.py projections.xlsx --adp adp.xlsx --dry-run
    python3 tools/import-projections.py projections.xlsx --adp adp.xlsx

The projections workbook has two sheets, found by name:

    skaters   NAME, POS, TEAM, ADP, ..., G, A, PTS
    goalies   NAME, POS, TEAM, AGE, ADP, W, SO

The ADP workbook's first sheet has NAME and AVG ADP (the mean of Yahoo and
Fantrax). Use it. The projections' own ADP is Yahoo's alone, which stops
around pick 153 and leaves a third of the pool with none. The bots price a
missing ADP as pick 360, so a 60-point forward with no ADP sank below
30-point players who had one. Without --adp the projections' ADP is used.

Columns are found by header, not position, so a reordered or widened sheet
still reads. "—" or "-" means no value.

What it computes, and why:

  points  skaters: PTS as given (goals + assists, the league's scoring).
          goalies: 2 x W + 3 x SO, the league's goalie scoring.
  vorp    points - BASELINE[position], the same baselines recompute-vorp.py
          uses, so a pool built here and one rebuilt there agree.
  tier    fixed point bands per position (TIERS below). They are the cut
          lines the previous pool used, so "tier 3" keeps its meaning across
          refreshes, and the scarcity rings on the board read the same.

Who is in the pool: every goalie; every skater with a projections-workbook ADP;
and skaters without one who project at or above FLOOR for their position. This
uses the projections' ADP even with --adp, because the ADP file lists nearly
every NHL skater and would pull in hundreds of fourth-liners. That keeps the board
at roughly the depth of a 14-team draft without a tail of fourth-liners.

Ids are the join key for every saved keeper, pick, queue entry and scouting
note, so a player already in the pool keeps his id exactly. Only a player who
is new gets a fresh one. The report lists who arrived and who left, and any
keeper in js/league.js that is missing from the result stops the write.

Source errata are corrected here rather than by hand after the fact; see
ERRATA.
"""

import json
import os
import re
import sys
import unicodedata

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FILE = os.path.join(ROOT, 'data', 'players.json')
LEAGUE_JS = os.path.join(ROOT, 'js', 'league.js')

BASELINE = {'F': 35.0, 'D': 25.5, 'G': 36.5}

# Lower bound of tiers 1..7 per position; anything below the last is tier 8.
TIERS = {
    'F': [100, 90, 80, 70, 61, 50, 42],
    'D': [80, 65, 55.5, 50, 45.5, 40, 31],
    'G': [72, 65, 61, 50, 41, 36, 32],
}

# Skaters with no ADP get in only at or above this projection.
FLOOR = {'F': 32.0, 'D': 22.5}

# The source lists the Canucks FORWARD Elias Pettersson as a D (his line is
# 20 G / 41 A over 80 games). The actual defenceman is "Elias Pettersson (D)".
ERRATA = {'Elias Pettersson': {'position': 'F'}}

DUMP = dict(separators=(',', ':'), ensure_ascii=False)


def fold(name):
    """Accent- and punctuation-free key, so 'Stützle' matches 'Stutzle'."""
    s = unicodedata.normalize('NFKD', name)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]', '', s.lower())


def slug(name):
    s = unicodedata.normalize('NFKD', name)
    s = ''.join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')


def num(v):
    if v is None:
        return None
    s = str(v).strip()
    if s in ('', '—', '-', 'N/A'):
        return None
    return float(s)


def position(pos):
    pos = str(pos).upper()
    if pos == 'G':
        return 'G'
    if pos == 'D':
        return 'D'
    return 'F'  # C, LW, RW and any combination of them


def read_sheet(wb, name):
    ws = (wb.worksheets[0] if name is None else
          next((s for s in wb.worksheets if s.title.strip().lower() == name), None))
    if ws is None:
        sys.exit('no "%s" sheet in the workbook (found: %s)' %
                 (name, ', '.join(s.title for s in wb.worksheets)))
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for i, r in enumerate(rows):
        if r and str(r[0] or '').strip().upper() == 'NAME':
            head = [str(c or '').strip().upper().rstrip('.') for c in r]
            out = []
            for body in rows[i + 1:]:
                if not body or not body[0]:
                    continue
                out.append({h: body[j] if j < len(body) else None
                            for j, h in enumerate(head) if h})
            return out
    sys.exit('no NAME header row in the "%s" sheet' % name)


def tier_for(pos, points):
    for i, cut in enumerate(TIERS[pos]):
        if points >= cut:
            return i + 1
    return len(TIERS[pos]) + 1


def keeper_ids():
    with open(LEAGUE_JS, encoding='utf-8') as fh:
        src = fh.read()
    block = src[src.index('keepers:'):src.index('trades:')]
    return re.findall(r"'([a-z0-9-]+)'", block)


def read_adp(path):
    import openpyxl
    rows = read_sheet(openpyxl.load_workbook(path, read_only=True, data_only=True), None)
    col = next((c for c in ('AVG ADP', 'ADP') if rows and c in rows[0]), None)
    if col is None:
        sys.exit('no AVG ADP or ADP column in %s' % path)
    return {fold(str(r['NAME'])): num(r.get(col)) for r in rows if r.get('NAME')}


def build(skaters, goalies, existing, adp_file=None):
    by_fold = {fold(p['name']): p['id'] for p in existing}
    taken = {p['id'] for p in existing}
    pool = []

    def add(name, pos, team, adp, points):
        if adp_file is not None:
            adp = adp_file.get(fold(name))
        pid = by_fold.get(fold(name))
        if pid is None:
            pid = slug(name)
            while pid in taken:
                pid += '-2'
            taken.add(pid)
        pool.append({
            'id': pid, 'name': name, 'position': pos, 'team': str(team).strip(),
            'tier': tier_for(pos, points),
            'vorp': round(points - BASELINE[pos], 2),
            'points': round(points, 2),
            'adp': adp,
        })

    for r in skaters:
        name = str(r['NAME']).strip()
        pos = ERRATA.get(name, {}).get('position', position(r['POS']))
        if pos == 'G':
            continue  # goalies come from their own sheet, which has W and SO
        pts = num(r.get('PTS'))
        if pts is None:
            continue
        adp = num(r.get('ADP'))
        if adp is None and pts < FLOOR[pos]:
            continue
        add(name, pos, r['TEAM'], adp, pts)

    for r in goalies:
        w, so = num(r.get('W')), num(r.get('SO'))
        if w is None or so is None:
            continue
        add(str(r['NAME']).strip(), 'G', r['TEAM'], num(r.get('ADP')), 2 * w + 3 * so)

    pool.sort(key=lambda p: (p['tier'], -p['vorp'], -p['points'], p['name']))
    return pool


def main():
    argv = sys.argv[1:]
    adp_path = None
    if '--adp' in argv:
        i = argv.index('--adp')
        adp_path = argv[i + 1]
        del argv[i:i + 2]
    args = [a for a in argv if not a.startswith('--')]
    if not args:
        sys.exit(__doc__)
    try:
        import openpyxl
    except ImportError:
        sys.exit('needs openpyxl:  pip install openpyxl')

    wb = openpyxl.load_workbook(args[0], read_only=True, data_only=True)
    skaters, goalies = read_sheet(wb, 'skaters'), read_sheet(wb, 'goalies')

    with open(FILE, encoding='utf-8') as fh:
        existing = json.load(fh)
    adp_file = read_adp(adp_path) if adp_path else None
    pool = build(skaters, goalies, existing, adp_file)

    old = {p['id']: p for p in existing}
    new = {p['id']: p for p in pool}
    counts = {pos: sum(1 for p in pool if p['position'] == pos) for pos in 'FDG'}
    print('%d players (%d F / %d D / %d G), %d with no ADP' % (
        len(pool), counts['F'], counts['D'], counts['G'],
        sum(1 for p in pool if p['adp'] is None)))
    for pos in 'FDG':
        per = [sum(1 for p in pool if p['position'] == pos and p['tier'] == t) for t in range(1, 9)]
        print('  %s tiers 1-8: %s' % (pos, ' '.join(str(n) for n in per)))

    arrived = sorted(new.keys() - old.keys())
    left = sorted(old.keys() - new.keys())
    print('new to the pool (%d): %s' % (len(arrived), ', '.join(new[i]['name'] for i in arrived)))
    print('dropped (%d): %s' % (len(left), ', '.join(old[i]['name'] for i in left)))
    moved = [(new[i]['name'], old[i]['team'], new[i]['team'])
             for i in new.keys() & old.keys() if old[i]['team'] != new[i]['team']]
    print('changed team (%d): %s' % (len(moved), ', '.join('%s %s>%s' % m for m in sorted(moved))))
    swapped = [(new[i]['name'], old[i]['position'], new[i]['position'])
               for i in new.keys() & old.keys() if old[i]['position'] != new[i]['position']]
    if swapped:
        print('CHANGED POSITION: %s' % swapped)

    missing = [k for k in keeper_ids() if k not in new]
    if missing:
        sys.exit('refusing to write: keepers missing from the new pool: %s' % ', '.join(missing))

    if '--dry-run' in sys.argv:
        print('dry run — nothing written')
        return
    with open(FILE, 'w', encoding='utf-8') as fh:
        fh.write(json.dumps(pool, **DUMP))
    print('wrote', os.path.relpath(FILE, ROOT))


if __name__ == '__main__':
    main()
